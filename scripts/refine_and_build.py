import os
import json
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

# 1. Register Fonts
pdfmetrics.registerFont(TTFont('YaHei', r'C:\Windows\Fonts\msyh.ttc', subfontIndex=0))
pdfmetrics.registerFont(TTFont('YaHeiBd', r'C:\Windows\Fonts\msyhbd.ttc', subfontIndex=0))
pdfmetrics.registerFont(TTFont('Malgun', r'C:\Windows\Fonts\malgun.ttf'))
pdfmetrics.registerFont(TTFont('MalgunBd', r'C:\Windows\Fonts\malgunbd.ttf'))

# POS translation map
pos_dict = {
    '名': '명사',
    '动': '동사',
    '形': '형용사',
    '副': '부사',
    '代': '대명사',
    '连': '접속사',
    '介': '개사',
    '量': '양사',
    '助': '조사',
    '叹': '감탄사',
    '数': '수사',
    '数量': '수량사',
    '后缀': '접미사'
}

def format_pos(pos_str):
    clean = pos_str.replace('\n', '').strip()
    ko = clean
    for k in sorted(pos_dict.keys(), key=lambda x: -len(x)):
        ko = ko.replace(k, pos_dict[k])
    ko = ko.replace('、', ', ')
    return clean, ko

def wrap_cjk_in_korean(text):
    t = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    # wrap CJK characters with YaHei font tag
    def replace_cjk(m):
        return f'<font name="YaHei">{m.group(0)}</font>'
    return re.sub(r'[\u4e00-\u9fff]+', replace_cjk, t)

def wrap_mixed_pos(clean_pos, ko_pos):
    return f'<font name="YaHei">{clean_pos}</font> <font name="Malgun">({ko_pos})</font>'

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont('Malgun', 8)
        self.setFillColor(colors.HexColor('#718096'))
        
        # Header (pages > 1)
        if self._pageNumber > 1:
            self.drawString(24, A4[1] - 20, "HSK 4 · Words 1001-2000 (HSK 3.0 한국어 대역 어휘집)")
            self.drawRightString(A4[0] - 24, A4[1] - 20, "studycli.org / HSK 표준대역")
            self.setStrokeColor(colors.HexColor('#E2E8F0'))
            self.setLineWidth(0.5)
            self.line(24, A4[1] - 24, A4[0] - 24, A4[1] - 24)

        # Footer
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(A4[0] - 24, 14, page_str)
        self.drawString(24, 14, "HSK 4급 어휘 목록 (1001-2000) · 한국어 대역판")
        self.setStrokeColor(colors.HexColor('#E2E8F0'))
        self.setLineWidth(0.5)
        self.line(24, 22, A4[0] - 24, 22)
        
        self.restoreState()

def build_pdf(items):
    pdf_filename = "HSK4_词汇表_韩文对照版.pdf"
    doc = SimpleDocTemplate(
        pdf_filename,
        pagesize=A4,
        leftMargin=24,
        rightMargin=24,
        topMargin=28,
        bottomMargin=28
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'MainTitle',
        fontName='YaHeiBd',
        fontSize=15,
        leading=18,
        alignment=1, # Center
        textColor=colors.HexColor('#1F4E79')
    )
    sub_style = ParagraphStyle(
        'SubTitle',
        fontName='Malgun',
        fontSize=8.5,
        leading=12,
        alignment=1, # Center
        textColor=colors.HexColor('#4A5568')
    )
    hdr_cell_style = ParagraphStyle(
        'HeaderCell',
        fontName='YaHeiBd',
        fontSize=8.5,
        leading=11,
        alignment=1, # Center
        textColor=colors.white
    )
    cell_center_yahei = ParagraphStyle(
        'CellCenterYahei',
        fontName='YaHei',
        fontSize=8,
        leading=10.5,
        alignment=1 # Center
    )
    cell_word_yahei = ParagraphStyle(
        'CellWordYahei',
        fontName='YaHeiBd',
        fontSize=8.5,
        leading=11,
        alignment=1,
        textColor=colors.HexColor('#1A365D')
    )
    cell_left_malgun = ParagraphStyle(
        'CellLeftMalgun',
        fontName='Malgun',
        fontSize=7.8,
        leading=10,
        alignment=0 # Left
    )
    cell_pos_style = ParagraphStyle(
        'CellPosStyle',
        fontName='Malgun',
        fontSize=7.5,
        leading=9.5,
        alignment=1 # Center
    )

    story = []
    story.append(Paragraph('<font name="YaHeiBd">HSK 4 · Words 1001-2000 · HSK 3.0</font> <font name="MalgunBd">(한국어 대역판)</font>', title_style))
    story.append(Spacer(1, 3))
    story.append(Paragraph('HSK 4급 어휘집 · 영문 해설을 표준 한국어 뜻으로 전면 개정 정리 (총 1,000단어)', sub_style))
    story.append(Spacer(1, 8))

    # Available width = 595.27 - 48 = 547.27
    col_widths = [36, 56, 72, 298, 85]

    table_data = [[
        Paragraph("No.", hdr_cell_style),
        Paragraph("汉字", hdr_cell_style),
        Paragraph("拼音", hdr_cell_style),
        Paragraph('<font name="YaHeiBd">韩文解释</font> <font name="MalgunBd">(한국어 뜻)</font>', hdr_cell_style),
        Paragraph('<font name="YaHeiBd">词性</font> <font name="MalgunBd">(품사)</font>', hdr_cell_style)
    ]]

    for it in items:
        styled_korean = wrap_cjk_in_korean(it["korean"])
        styled_pos = wrap_mixed_pos(it["pos_clean"], it["pos_ko"])
        row = [
            Paragraph(str(it["no"]), cell_center_yahei),
            Paragraph(it["word"], cell_word_yahei),
            Paragraph(it["pinyin"], cell_center_yahei),
            Paragraph(styled_korean, cell_left_malgun),
            Paragraph(styled_pos, cell_pos_style)
        ]
        table_data.append(row)

    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E79')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.2),
        ('TOPPADDING', (0, 0), (-1, -1), 2.2),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#1F4E79')),
    ]

    for r in range(1, len(table_data)):
        if r % 2 == 0:
            t_style.append(('BACKGROUND', (0, r), (-1, r), colors.HexColor('#F8FAFC')))
        else:
            t_style.append(('BACKGROUND', (0, r), (-1, r), colors.white))

    t.setStyle(TableStyle(t_style))
    story.append(t)

    doc.build(story, canvasmaker=NumberedCanvas)
    print("Saved HSK4_词汇表_韩文对照版.pdf")

def main():
    with open('HSK4_词汇表_韩文版.json', 'r', encoding='utf-8') as f:
        items = json.load(f)

    curations = {
        1072: ("厂", "공장"),
        1073: ("场", "마당, 장소; (경기·공연·일 등의 횟수를 세는) 판, 회, 차례"),
        1078: ("乘1", "(차·배·비행기 등을) 타다, 탑승하다; 곱하다"),
        1079: ("成功", "성공하다; 성공"),
        1147: ("等2", "등, 따위; 등급"),
        1220: ("干", "마르다, 건조하다; 유능하다"),
        1224: ("感", "느끼다, 감각; (접미사: ~감, 느낌)"),
        1430: ("克", "그램 (g); 이기다, 극복하다"),
        1473: ("两2", "냥 (중국 무게 단위, 1냥=50g)"),
        1496: ("满", "가득 차다, 차다; 만족하다; 만주족"),
        1538: ("弄", "하다, 만들다, 처리하다; 골목"),
        1543: ("排", "줄, 열; 줄을 서다, 배열하다; 줄 (양사)"),
        1544: ("牌", "패, 카드, 번호판, 상표, 간판"),
        1646: ("生1", "낳다, 태어나다; 생기다, 자라다"),
        1647: ("生2", "날것의, 설익은; 낯설다, 서투르다"),
        1650: ("省1", "성 (중국 행정 구역 단위)"),
        1718: ("弹", "(악기를) 연주하다, 타다; 튕기다"),
        1743: ("通", "통하다, 뚫리다; 정통하다, 잘 알다; 통 (편지·전화 양사)"),
        1771: ("为1", "~이 되다, ~을 하다; ~으로 삼다, 간주하다"),
        1938: ("之1", "~의 (문어체 '的'); 그것, 그 사람"),
        1959: ("种", "(씨앗·나무 등을) 심다"),
        1960: ("重", "무겁다, 중요하다; 심하다"),
        1970: ("转", "돌다, 회전하다; 방향을 바꾸다"),
        1972: ("转机1", "(비행기를) 갈아타다, 환승하다"),
        1996: ("作为1", "~로 여기다, 삼다; ~로서 (신분·자격)"),
    }

    for it in items:
        clean_pos, ko_pos = format_pos(it['pos'])
        it['pos_clean'] = clean_pos
        it['pos_ko'] = ko_pos
        it['pos_display'] = f"{clean_pos} ({ko_pos})"
        
        if it['no'] in curations:
            it['korean'] = curations[it['no']][1]

    # 1. JSON
    with open('HSK4_词汇表_韩文对照版.json', 'w', encoding='utf-8') as f:
        json.dump(items, f, ensure_ascii=False, indent=2)
    print("1. Saved HSK4_词汇表_韩文对照版.json")

    # 2. CSV
    with open('HSK4_词汇表_韩文对照版.csv', 'w', encoding='utf-8-sig') as f:
        f.write("序号,汉字,拼音,韩文解释,词性,韩文词性\n")
        for it in items:
            ko = f'"{it["korean"]}"' if ',' in it["korean"] else it["korean"]
            f.write(f'{it["no"]},{it["word"]},{it["pinyin"]},{ko},{it["pos_clean"]},{it["pos_ko"]}\n')
    print("2. Saved HSK4_词汇表_韩文对照版.csv")

    # 3. Markdown
    with open('HSK4_词汇表_韩文对照版.md', 'w', encoding='utf-8') as f:
        f.write("# HSK 4 级词汇表（1001 - 2000）· 韩文对照版\n\n")
        f.write("> **说明**：本词汇表由原 HSK 4 级英文对照表（The Chinese Language Institute, studycli.org）整理升级，将原有英文解释全面转换为标准、精准的韩文释义（한국어 뜻），并配以对应韩文词性，全表共收录 1,000 词（No. 1001 ~ 2000）。\n\n")
        f.write("| 序号 (No.) | 汉字 (단어) | 拼音 (병음) | 韩文解释 (한국어 뜻) | 词性 (품사) |\n")
        f.write("| :---: | :---: | :---: | :--- | :---: |\n")
        for it in items:
            korean = it["korean"].replace("|", "/")
            f.write(f'| {it["no"]} | {it["word"]} | {it["pinyin"]} | {korean} | {it["pos_display"]} |\n')
    print("3. Saved HSK4_词汇表_韩文对照版.md")

    # 4. Excel (.xlsx)
    wb = Workbook()
    ws = wb.active
    ws.title = "HSK 4 词汇表 (韩文版)"
    ws.views.sheetView[0].showGridLines = True

    # Title
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value = "HSK 4 级大纲词汇表 (1001 - 2000) · 韩文对照版"
    title_cell.font = Font(name="Microsoft YaHei", size=15, bold=True, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Subtitle
    ws.merge_cells("A2:E2")
    sub_cell = ws["A2"]
    sub_cell.value = "HSK 3.0 표준 어휘집 | 영문 해설을 한국어 뜻으로 전면 개정 | 총 1,000단어 수록"
    sub_cell.font = Font(name="Malgun Gothic", size=10, italic=True, color="595959")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    ws.row_dimensions[3].height = 6

    # Headers on row 4
    headers = ["序号 (No.)", "汉字 (단어)", "拼音 (병음)", "韩文解释 (한국어 뜻)", "词性 (품사)"]
    ws.row_dimensions[4].height = 26

    header_font = Font(name="Malgun Gothic", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_border = Border(
        left=Side(style='thin', color='B0C4DE'),
        right=Side(style='thin', color='B0C4DE'),
        top=Side(style='medium', color='1F4E79'),
        bottom=Side(style='medium', color='1F4E79')
    )

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border

    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra = PatternFill(start_color="F5F8FC", end_color="F5F8FC", fill_type="solid")
    
    font_regular = Font(name="Malgun Gothic", size=10)
    font_bold_word = Font(name="Microsoft YaHei", size=11, bold=True, color="002060")
    font_pinyin = Font(name="Segoe UI", size=10, color="333333")
    font_pos = Font(name="Malgun Gothic", size=9.5, color="404040")

    for r_idx, item in enumerate(items, start=5):
        row_fill = fill_zebra if r_idx % 2 == 0 else fill_white
        ws.row_dimensions[r_idx].height = 22

        c_no = ws.cell(row=r_idx, column=1, value=item["no"])
        c_no.alignment = Alignment(horizontal="center", vertical="center")
        c_no.font = font_regular
        c_no.fill = row_fill
        c_no.border = thin_border

        c_word = ws.cell(row=r_idx, column=2, value=item["word"])
        c_word.alignment = Alignment(horizontal="center", vertical="center")
        c_word.font = font_bold_word
        c_word.fill = row_fill
        c_word.border = thin_border

        c_py = ws.cell(row=r_idx, column=3, value=item["pinyin"])
        c_py.alignment = Alignment(horizontal="center", vertical="center")
        c_py.font = font_pinyin
        c_py.fill = row_fill
        c_py.border = thin_border

        c_ko = ws.cell(row=r_idx, column=4, value=item["korean"])
        c_ko.alignment = Alignment(horizontal="left", vertical="center")
        c_ko.font = font_regular
        c_ko.fill = row_fill
        c_ko.border = thin_border

        c_pos = ws.cell(row=r_idx, column=5, value=item["pos_display"])
        c_pos.alignment = Alignment(horizontal="center", vertical="center")
        c_pos.font = font_pos
        c_pos.fill = row_fill
        c_pos.border = thin_border

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 52
    ws.column_dimensions['E'].width = 24
    ws.freeze_panes = "A5"

    wb.save('HSK4_词汇表_韩文对照版.xlsx')
    print("4. Saved HSK4_词汇表_韩文对照版.xlsx")

    # 5. Word (.docx)
    doc = Document()
    doc.add_heading("HSK 4 级大纲词汇表 (1001 - 2000) · 韩文对照版", level=1)
    p = doc.add_paragraph("HSK 3.0 표준 어휘집 | 영문 해설을 한국어 뜻으로 전면 개정 | 총 1,000단어 수록")
    p.italic = True

    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    hdr_titles = ["序号", "汉字", "拼音", "韩文解释", "词性 (품사)"]
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.name = "Malgun Gothic"

    col_widths = [Inches(0.7), Inches(0.9), Inches(1.1), Inches(3.4), Inches(1.3)]
    for i, col in enumerate(table.columns):
        for cell in col.cells:
            cell.width = col_widths[i]

    for it in items:
        row_cells = table.add_row().cells
        row_cells[0].text = str(it["no"])
        row_cells[1].text = it["word"]
        row_cells[2].text = it["pinyin"]
        row_cells[3].text = it["korean"]
        row_cells[4].text = it["pos_display"]
        for i, cell in enumerate(row_cells):
            cell.width = col_widths[i]

    doc.save('HSK4_词汇表_韩文对照版.docx')
    print("5. Saved HSK4_词汇表_韩文对照版.docx")

    # 6. PDF
    build_pdf(items)

if __name__ == '__main__':
    main()
